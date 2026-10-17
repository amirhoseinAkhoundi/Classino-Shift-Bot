"""خروجی Excel آمار."""
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from analysis import fmt_entry, hours, sort_entries, stats
from schedule_parser import DAYS


def _sheet(wb, title, header, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    return ws


def build_xlsx(entries, path):
    studios, admins = stats(entries)
    wb = Workbook()
    _sheet(wb, "استودیوها", ["استودیو", "ساعت", "تعداد کلاس"],
           [[n, float(hours(m)), c] for n, m, c in studios], first=True)
    _sheet(wb, "ادمین‌ها", ["نام", "ساعت کلاس", "ساعت نظارت", "ساعت کمکی", "جمع ساعت"],
           [[a["person"], float(hours(a["class"])), float(hours(a["nazer"])), float(hours(a["komaki"])),
             float(hours(a["total"]))] for a in admins])
    days_present = [d for d in DAYS if any(e["day"] == d for e in entries)]
    per = defaultdict(lambda: defaultdict(int))
    for e in entries:
        if e["person"]:
            per[e["person"]][e["day"]] += e["end_min"] - e["start_min"]
    _sheet(wb, "ساعت هر روز", ["نام"] + days_present,
           [[p] + [float(hours(per[p].get(d, 0))) for d in days_present] for p in sorted(per)])
    _sheet(wb, "جزئیات", ["روز", "نوع", "شرح", "فرد"],
           [[e["day"], e["kind"], fmt_entry(e), e["person"] or ""] for e in sort_entries(entries)])
    wb.save(path)
    return path


def build_studios_xlsx(entries, path):
    studios, _ = stats(entries)
    wb = Workbook()
    _sheet(wb, "استودیوها", ["استودیو", "مجموع ساعت استفاده", "تعداد کلاس"],
           [[name, float(hours(minutes)), count] for name, minutes, count in studios], first=True)
    wb.save(path)
    return path


def build_admins_xlsx(entries, path):
    _, admins = stats(entries)
    wb = Workbook()
    _sheet(wb, "ادمین‌ها", ["نام", "ساعت کلاس", "ساعت نظارت", "ساعت کمکی", "جمع ساعت"],
           [[a["person"], float(hours(a["class"])), float(hours(a["nazer"])), float(hours(a["komaki"])),
             float(hours(a["total"]))] for a in admins], first=True)
    wb.save(path)
    return path
